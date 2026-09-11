"""Render complete frame/bin plot families for z-binned strain results.

The functions in this module reproduce the specialized batch plots emitted by
the original standalone analyses.  Numerical analysis remains in
``reaxkit.analysis``; this module only consumes result tables.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Sequence

import numpy as np
import pandas as pd

from reaxkit.analysis.stress_strain.z_binned_deformation_gradient_strain import (
    ENGINEERING_SHEAR_COLUMNS,
    NORMAL_STRAIN_COLUMNS,
)
from reaxkit.core.platform.paths import io_path

CHANGE_COLUMNS = ("span_change_x", "span_change_y", "span_change_z")
STRAIN_COLUMNS = ("strain_xx", "strain_yy", "strain_zz")
LEFT_NORMAL_COLUMNS = ("strain_xx", "strain_yy")
RIGHT_NORMAL_COLUMNS = ("strain_zz",)
CORRELATION_PAIRS = (
    ("span_change_x", "span_change_y"),
    ("span_change_x", "span_change_z"),
    ("span_change_y", "span_change_z"),
)
PLOT_LABELS = {
    "span_change_x": r"$\mathrm{span\ change}_{x}$",
    "span_change_y": r"$\mathrm{span\ change}_{y}$",
    "span_change_z": r"$\mathrm{span\ change}_{z}$",
    "strain_xx": r"$\mathrm{strain}_{xx}$",
    "strain_yy": r"$\mathrm{strain}_{yy}$",
    "strain_zz": r"$\mathrm{strain}_{zz}$",
    "gamma_xy": r"$\gamma_{xy}$",
    "gamma_xz": r"$\gamma_{xz}$",
    "gamma_yz": r"$\gamma_{yz}$",
}


def _plotting_module():
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:  # pragma: no cover - depends on optional extras
        raise ImportError(
            "Plot output requires matplotlib. Install the ReaxKit plot extra: reaxkit[plot]."
        ) from exc
    return plt


def _validate_table(table: pd.DataFrame, columns: Sequence[str]) -> None:
    required = {"frame", "iter", "bin_number", "bin_mean_z", *columns}
    missing = sorted(required.difference(table.columns))
    if table.empty:
        raise ValueError("Cannot generate batch plots from an empty result table.")
    if missing:
        raise ValueError(f"Result table is missing plot columns: {', '.join(missing)}")


def _padded_limits(
    table: pd.DataFrame,
    columns: Sequence[str],
    *,
    scale: float = 1.0,
    symmetric: bool = False,
) -> tuple[float, float] | None:
    values = table[list(columns)].to_numpy(dtype=float) * scale
    finite = values[np.isfinite(values)]
    if finite.size == 0:
        return None
    if symmetric:
        maximum = float(np.max(np.abs(finite)))
        extent = 1.0e-6 if maximum <= np.finfo(float).eps else 1.05 * maximum
        return -extent, extent
    lower = min(float(finite.min()), 0.0)
    upper = max(float(finite.max()), 0.0)
    span = upper - lower
    padding = (
        0.05 * span
        if span > np.finfo(float).eps * max(1.0, abs(lower), abs(upper))
        else max(abs(lower), abs(upper), 1.0) * 1.0e-6
    )
    return lower - padding, upper + padding


def _correlations(subset: pd.DataFrame) -> dict[str, tuple[float | None, int]]:
    output: dict[str, tuple[float | None, int]] = {}
    for first, second in CORRELATION_PAIRS:
        first_values = subset[first].to_numpy(dtype=float)
        second_values = subset[second].to_numpy(dtype=float)
        valid = np.isfinite(first_values) & np.isfinite(second_values)
        count = int(valid.sum())
        correlation = None
        if count >= 2:
            x = first_values[valid] - float(first_values[valid].mean())
            y = second_values[valid] - float(second_values[valid].mean())
            denominator = float(np.linalg.norm(x) * np.linalg.norm(y))
            if denominator > 0.0:
                correlation = float(np.clip(np.dot(x, y) / denominator, -1.0, 1.0))
        output[f"{first}-{second}"] = (correlation, count)
    return output


def _correlation_annotation(subset: pd.DataFrame) -> str:
    lines = ["Pearson component correlation"]
    for pair, (value, count) in _correlations(subset).items():
        lines.append(f"{pair}: {'n/a' if value is None else f'{value:+.3f}'} (n={count})")
    return "\n".join(lines)


def _correlation_table(
    table: pd.DataFrame, *, group_by: str, group_values: Sequence[int]
) -> pd.DataFrame:
    group_column = "frame" if group_by == "frame" else "bin_number"
    identifiers = (
        ["frame", "iter", "number_of_z_bins"]
        if group_by == "frame"
        else ["bin_number", "number_of_frames"]
    )
    metrics = [
        column
        for first, second in CORRELATION_PAIRS
        for column in (f"pearson_r_{first}_{second}", f"n_{first}_{second}")
    ]
    rows = []
    for value in group_values:
        subset = table[table[group_column].astype(int) == value]
        row = (
            {"frame": value, "iter": int(subset["iter"].iloc[0]), "number_of_z_bins": len(subset)}
            if group_by == "frame"
            else {"bin_number": value, "number_of_frames": len(subset)}
        )
        for pair, (correlation, count) in _correlations(subset).items():
            first, second = pair.split("-")
            row[f"pearson_r_{first}_{second}"] = correlation
            row[f"n_{first}_{second}"] = count
        rows.append(row)
    return pd.DataFrame(rows, columns=[*identifiers, *metrics])


def _excel_value(value):
    if value is None or (isinstance(value, (float, np.floating)) and not np.isfinite(value)):
        return None
    return value.item() if isinstance(value, np.generic) else value


def _write_correlations(table: pd.DataFrame, destination: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:  # pragma: no cover - depends on optional extras
        raise ImportError(
            "Correlation workbooks require openpyxl. Install the ReaxKit io extra: reaxkit[io]."
        ) from exc
    io_path(destination.parent).mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Span_Change_Correlations"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.append(list(table.columns))
    for row in table.itertuples(index=False, name=None):
        sheet.append([_excel_value(value) for value in row])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for number, name in enumerate(table.columns, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = max(12, min(26, len(name) + 2))
    reference = f"A1:{get_column_letter(len(table.columns))}{sheet.max_row}"
    excel_table = Table(displayName="SpanChangeCorrelations", ref=reference)
    excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(excel_table)
    temporary = destination.with_name(f".{destination.stem}.tmp.xlsx")
    workbook.save(io_path(temporary))
    workbook.close()
    os.replace(io_path(temporary), io_path(destination))


def _plot_family(
    table: pd.DataFrame,
    output_root: Path,
    *,
    group_by: str,
    family: str,
    columns: Sequence[str],
    scale: float,
    dpi: int,
    plot_every: int,
    y_scale: str,
    primary_limits: tuple[float, float] | None,
    secondary_limits: tuple[float, float] | None = None,
    correlations: bool = False,
) -> list[Path]:
    plt = _plotting_module()
    group_column = "frame" if group_by == "frame" else "bin_number"
    folder_name = f"{family}_vs_z_by_frame" if group_by == "frame" else f"{family}_vs_frame_by_bin"
    folder = output_root / folder_name
    io_path(folder).mkdir(parents=True, exist_ok=True)
    group_values = sorted(table[group_column].astype(int).unique())[::plot_every]
    colors = ("#C43C39", "#2E8B57", "#3569B0")
    outputs: list[Path] = []
    for group_value in group_values:
        subset = table[table[group_column].astype(int) == group_value].sort_values(
            "bin_number" if group_by == "frame" else "frame"
        )
        x_values = subset["bin_mean_z"] if group_by == "frame" else subset["frame"]
        title_suffix = (
            f"frame {group_value}, iteration {int(subset['iter'].iloc[0])}"
            if group_by == "frame"
            else f"bin {group_value}"
        )
        figure, axis = plt.subplots(figsize=(9.5, 5.8), constrained_layout=True)
        secondary_axis = axis.twinx() if family in {"strain", "normal_strain"} else None
        lines = []
        for column, color in zip(columns, colors):
            plot_axis = secondary_axis if column in RIGHT_NORMAL_COLUMNS and secondary_axis is not None else axis
            (line,) = plot_axis.plot(
                x_values, subset[column] * scale, color=color, linewidth=1.7,
                marker="o", markersize=3.5, label=PLOT_LABELS[column],
            )
            lines.append(line)
        axis.axhline(0.0, color="0.35", linewidth=0.9, linestyle="--")
        axis.set_xlabel("z-bin midpoint (angstrom)" if group_by == "frame" else "Frame index")
        if secondary_axis is None:
            axis.set_ylabel("Strain (%)" if scale == 100.0 else "Span change (angstrom)")
            if y_scale == "global" and primary_limits is not None:
                axis.set_ylim(*primary_limits)
        else:
            axis.set_ylabel(r"$\mathrm{strain}_{xx}$, $\mathrm{strain}_{yy}$ (%)")
            secondary_axis.set_ylabel(r"$\mathrm{strain}_{zz}$ (%)", color=colors[2])
            secondary_axis.tick_params(axis="y", colors=colors[2])
            secondary_axis.spines["right"].set_color(colors[2])
            left = primary_limits if y_scale == "global" else _padded_limits(subset, LEFT_NORMAL_COLUMNS, scale=100.0, symmetric=True)
            right = secondary_limits if y_scale == "global" else _padded_limits(subset, RIGHT_NORMAL_COLUMNS, scale=100.0, symmetric=True)
            if left is not None:
                axis.set_ylim(*left)
            if right is not None:
                secondary_axis.set_ylim(*right)
        scale_label = "global y-scale" if y_scale == "global" else "frame-wise y-scale"
        axis.set_title(f"{family.replace('_', ' ').title()}; {title_suffix}; {scale_label}")
        axis.grid(True, alpha=0.25)
        axis.legend(handles=lines, frameon=True)
        if correlations:
            axis.text(
                0.015, 0.985, _correlation_annotation(subset), transform=axis.transAxes,
                ha="left", va="top", fontsize=8.5, linespacing=1.25,
                bbox={"boxstyle": "round,pad=0.4", "facecolor": "white", "edgecolor": "0.7", "alpha": 0.9},
            )
        destination = folder / f"{folder_name}_{group_column}_{group_value:06d}.png"
        figure.savefig(io_path(destination), dpi=dpi, bbox_inches="tight")
        plt.close(figure)
        outputs.append(destination)
    if correlations:
        _write_correlations(
            _correlation_table(table, group_by=group_by, group_values=group_values),
            folder / "correlations.xlsx",
        )
    return outputs


def generate_top_bottom_plots(
    table: pd.DataFrame,
    output_root: str | Path,
    *,
    dpi: int = 180,
    plot_every: int = 1,
    y_scale: str = "global",
) -> list[Path]:
    """Generate all frame/bin span-change and normal-strain plot families.

    Parameters
    ----------
    table : pd.DataFrame
        Top/bottom analyzer result table.
    output_root : str or Path
        Parent directory for the four plot-family directories.
    dpi : int
        PNG resolution.
    plot_every : int
        Render every Nth frame; bin histories are always complete.
    y_scale : str
        ``global`` shares family-wide limits; ``frame`` autoscales each plot.

    Returns
    -------
    list[Path]
        Generated PNG paths. Correlation workbooks are also written beside
        both span-change families.

    Examples
    --------
    ``generate_top_bottom_plots(result.table, "plots", y_scale="global")``
    writes four plot-family directories with family-wide axis limits.
    """
    _validate_table(table, (*CHANGE_COLUMNS, *STRAIN_COLUMNS))
    if dpi < 1 or plot_every < 1:
        raise ValueError("dpi and plot_every must be positive.")
    if y_scale not in {"global", "frame"}:
        raise ValueError("y_scale must be 'global' or 'frame'.")
    root = Path(output_root).resolve()
    span_limits = _padded_limits(table, CHANGE_COLUMNS) if y_scale == "global" else None
    left_limits = _padded_limits(table, LEFT_NORMAL_COLUMNS, scale=100.0, symmetric=True) if y_scale == "global" else None
    right_limits = _padded_limits(table, RIGHT_NORMAL_COLUMNS, scale=100.0, symmetric=True) if y_scale == "global" else None
    outputs = []
    for group_by, family, columns, scale in (
        ("frame", "span_change", CHANGE_COLUMNS, 1.0),
        ("frame", "strain", STRAIN_COLUMNS, 100.0),
        ("bin", "span_change", CHANGE_COLUMNS, 1.0),
        ("bin", "strain", STRAIN_COLUMNS, 100.0),
    ):
        outputs.extend(_plot_family(
            table, root, group_by=group_by, family=family, columns=columns, scale=scale,
            dpi=dpi, plot_every=plot_every if group_by == "frame" else 1,
            y_scale=y_scale, primary_limits=left_limits if family == "strain" else span_limits,
            secondary_limits=right_limits if family == "strain" else None,
            correlations=family == "span_change",
        ))
    return outputs


def generate_deformation_gradient_plots(
    table: pd.DataFrame,
    output_root: str | Path,
    *,
    dpi: int = 180,
    plot_every: int = 1,
    y_scale: str = "global",
) -> list[Path]:
    """Generate complete normal- and engineering-shear plot families.

    Parameters
    ----------
    table : pd.DataFrame
        Deformation-gradient analyzer result table.
    output_root : str or Path
        Parent directory for the four plot-family directories.
    dpi : int
        PNG resolution.
    plot_every : int
        Render every Nth frame; bin histories are always complete.
    y_scale : str
        ``global`` shares family-wide limits; ``frame`` autoscales each plot.

    Returns
    -------
    list[Path]
        Generated PNG paths.

    Examples
    --------
    ``generate_deformation_gradient_plots(result.table, "plots")`` writes
    normal-strain and engineering-shear families by frame and bin.
    """
    _validate_table(table, (*NORMAL_STRAIN_COLUMNS, *ENGINEERING_SHEAR_COLUMNS))
    if dpi < 1 or plot_every < 1:
        raise ValueError("dpi and plot_every must be positive.")
    if y_scale not in {"global", "frame"}:
        raise ValueError("y_scale must be 'global' or 'frame'.")
    root = Path(output_root).resolve()
    left_limits = _padded_limits(table, LEFT_NORMAL_COLUMNS, scale=100.0, symmetric=True) if y_scale == "global" else None
    right_limits = _padded_limits(table, RIGHT_NORMAL_COLUMNS, scale=100.0, symmetric=True) if y_scale == "global" else None
    shear_limits = _padded_limits(table, ENGINEERING_SHEAR_COLUMNS, scale=100.0) if y_scale == "global" else None
    outputs = []
    for family, columns in (
        ("normal_strain", NORMAL_STRAIN_COLUMNS),
        ("engineering_shear", ENGINEERING_SHEAR_COLUMNS),
    ):
        for group_by in ("frame", "bin"):
            outputs.extend(_plot_family(
                table, root, group_by=group_by, family=family, columns=columns, scale=100.0,
                dpi=dpi, plot_every=plot_every if group_by == "frame" else 1,
                y_scale=y_scale, primary_limits=left_limits if family == "normal_strain" else shear_limits,
                secondary_limits=right_limits if family == "normal_strain" else None,
            ))
    return outputs
