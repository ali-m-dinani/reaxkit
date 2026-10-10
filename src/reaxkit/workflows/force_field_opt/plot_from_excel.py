"""Render user-grouped Excel data with the shared ReaxKit plotting presets."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from reaxkit.core.platform.paths import io_path
from reaxkit.presentation.plot import plot as render_plot
from reaxkit.presentation.plot_styles import add_plot_style_argument
from reaxkit.presentation.color_styles import add_color_style_argument, resolve_color_style
from reaxkit.workflows.file_tools.ffield_workflow import QM_PLOT_COLOR, REAXFF_PLOT_COLOR

ALL_COMMANDS = ("plot-from-excel",)
ALL_LEGACY_COMMANDS = ()
SHEETS = ("bar_plots", "line_plots", "EOS_plots")
DATA_HEADERS = ("Figure number", "Line in fort.99", "Identifier", "X value", "ReaxFF value", "QM value")
SETTINGS_HEADERS = ("Figure number", "Title", "X-axis label", "Y-axis label")


def _blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def _number(value, location, *, integer=False):
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{location}: expected a finite number, got {value!r}.") from None
    if isinstance(value, bool) or not math.isfinite(number):
        raise ValueError(f"{location}: expected a finite number, got {value!r}.")
    if integer and (number < 1 or not number.is_integer()):
        raise ValueError(f"{location}: figure number must be a positive integer.")
    return int(number) if integer else number


def read_figures(path):
    """Validate all populated sheets before returning per-figure data and settings."""
    workbook = load_workbook(io_path(path), data_only=True)
    formulas = load_workbook(io_path(path), data_only=False)
    figures = []
    try:
        found = [name for name in SHEETS if name in workbook.sheetnames]
        if not found:
            raise ValueError(f"Workbook must contain at least one of: {', '.join(SHEETS)}.")
        for name in found:
            sheet = workbook[name]

            def value_at(row, column):
                cell = sheet.cell(row, column)
                original = formulas[name].cell(row, column)
                location = f"{name}!{cell.coordinate}"
                if original.data_type == "f" and cell.value is None:
                    raise ValueError(f"{location}: formula has no cached result. Recalculate and save in Excel first.")
                if cell.data_type == "e":
                    raise ValueError(f"{location}: Excel error {cell.value}.")
                return cell.value

            if all(_blank(cell.value) for row in sheet for cell in row):
                continue
            headers = tuple(str(value_at(1, column) or "").strip().lower() for column in range(1, 7))
            if headers != tuple(header.lower() for header in DATA_HEADERS):
                raise ValueError(f"{name}!A1:F1 must contain: {', '.join(DATA_HEADERS)}. Use the updated template.")
            settings_headers = tuple(str(value_at(1, column) or "").strip().lower() for column in range(8, 12))
            if settings_headers != tuple(header.lower() for header in SETTINGS_HEADERS):
                raise ValueError(f"{name}!H1:K1 must contain: {', '.join(SETTINGS_HEADERS)}.")
            settings = {}
            groups = {}
            for row in range(2, sheet.max_row+1):
                data = [value_at(row, column) for column in range(1, 7)]
                if not all(_blank(value) for value in data):
                    figure = _number(data[0], f"{name}!A{row}", integer=True)
                    if _blank(data[2]):
                        raise ValueError(f"{name}!C{row}: Identifier is required.")
                    coordinate = None if _blank(data[3]) else _number(data[3], f"{name}!D{row}")
                    if name == "EOS_plots" and coordinate is None:
                        raise ValueError(f"{name}!D{row}: EOS requires an X value (volume).")
                    groups.setdefault(figure, []).append({
                        "row": row, "line_in_fort99": data[1], "identifier": str(data[2]).strip(),
                        "x": coordinate, "reaxff": _number(data[4], f"{name}!E{row}"),
                        "qm": _number(data[5], f"{name}!F{row}"),
                    })
                values = [value_at(row, column) for column in range(8, 12)]
                if not all(_blank(value) for value in values):
                    figure = _number(values[0], f"{name}!H{row}", integer=True)
                    if figure in settings:
                        raise ValueError(f"{name}!H{row}: duplicate settings for figure {figure}.")
                    if any(_blank(value) for value in values[2:]):
                        raise ValueError(f"{name}!J{row}: both axis labels (including units) are required.")
                    settings[figure] = {"title": str(values[1] or ""), "xlabel": str(values[2]), "ylabel": str(values[3])}
            if set(settings) != set(groups):
                raise ValueError(f"{name}: settings figure numbers must match data; missing {sorted(set(groups)-set(settings))}, unused {sorted(set(settings)-set(groups))}.")
            for figure, rows in sorted(groups.items()):
                categorical = name == "line_plots" and all(row["x"] is None for row in rows)
                if name != "bar_plots":
                    if len(rows) < 2:
                        raise ValueError(f"{name} figure {figure}: curves require at least two points.")
                    if not categorical:
                        if any(row["x"] is None for row in rows):
                            raise ValueError(f"{name} figure {figure}: provide X values for every point, or leave all blank for categorical lines.")
                        rows = sorted(rows, key=lambda row: row["x"])
                figures.append({"sheet": name, "figure": figure, "rows": rows,
                                "categorical": categorical, **settings[figure]})
        if not figures:
            raise ValueError("Workbook contains no figure data.")
        return figures
    finally:
        workbook.close()
        formulas.close()


def figure_payload(figure, *, plot_style, figsize):
    """Build the same paired-series payloads used by force-field result plots."""
    rows = figure["rows"]
    bar = figure["sheet"] == "bar_plots"
    coordinates = list(range(len(rows))) if figure["categorical"] else [row["x"] for row in rows]
    payload = {"plot_type": "grouped_bar_plot" if bar else "single_plot",
               "plot_style": plot_style, "figsize": figsize, "legend": True,
               "grid": False, "title": figure["title"], "xlabel": figure["xlabel"],
               "ylabel": figure["ylabel"], "series": []}
    for field, label, color in (("reaxff", "ReaxFF", REAXFF_PLOT_COLOR), ("qm", "QM value", QM_PLOT_COLOR)):
        series = {"label": label, "color": color}
        values = [row[field] for row in rows]
        if bar:
            series["values"] = values
        else:
            series.update(x=coordinates, y=values, marker="o")
        payload["series"].append(series)
    if bar:
        long_labels = max(len(row["identifier"]) for row in rows) > 20
        payload.update(labels=[row["identifier"] for row in rows], group_width=0.48,
                       minimum_category_slots=3, label_rotation=45 if long_labels else 0,
                       label_horizontal_alignment="right" if long_labels else "center",
                       label_multialignment="left")
    elif figure["categorical"]:
        payload.update(xticks=coordinates, xticklabels=[row["identifier"] for row in rows],
                       xtick_rotation=20)
    return payload


def build_parser(parser, *, command=None):
    """Document workbook plotting inputs, scope, and manuscript export examples."""
    parser.formatter_class = argparse.RawTextHelpFormatter
    parser.description = (
        "Render Excel data as figures grouped independently by sheet and figure number.\n\n"
        "Use the figure-generator workbook copied by get_ffield_opt_plots to select and group training data.\n"
        "Supported sheets are bar_plots, line_plots, and EOS_plots; rendering uses shared ReaxKit styles.\n"
        "Set data in A:F and figure titles/axis labels with units in H:K.\n\n"
        "Scope: reads entered values and saved formula results, not fort.99 or GEO files.\n"
        "Recalculate and save formulas in Excel first. No simulation, EOS fit, sign change,\n"
        "or energy normalization is performed. Numeric curves sort by X; categorical lines\n"
        "use worksheet order when all X values are blank. EOS requires numeric volume values.\n\n"
        "Examples:\n"
        "  1. Render grouped bars, line profiles, and EOS curves from the prepared template:\n"
        "       reaxkit plot-from-excel --input template_ffield_opt_figure_generator.xlsx\n\n"
        "  2. Export manuscript figures with the same style as automatic ReaxKit plots:\n"
        "       reaxkit plot-from-excel --input PtZn_training_figures.xlsx --output manuscript_figures --plot-style publication-bold\n\n"
        "  3. Set panel dimensions and export raster and vector formats:\n"
        "       reaxkit plot-from-excel --input PtZn_training_figures.xlsx --width 8 --height 5.2 --formats png svg pdf\n\n"
        "  4. Regenerate existing figures after editing the workbook:\n"
        "       reaxkit plot-from-excel --input PtZn_training_figures.xlsx --output manuscript_figures --overwrite"
    )
    parser.add_argument("--input", required=True, help="Excel workbook containing figure data and settings. Example: --input PtZn_training_figures.xlsx, which reads the prepared training-data sheets.")
    parser.add_argument("--output", default=None, help="Directory for figures and their source-row manifest (default: excel_figures beside the input workbook). Example: --output manuscript_figures, which writes separate subfolders for each plot sheet in the specified directory.")
    parser.add_argument("--formats", nargs="+", choices=("png", "svg", "pdf"), default=["png", "svg"], help="Image formats exported for every figure. Example: --formats png svg pdf, which saves raster PNG and vector SVG/PDF versions.")
    parser.add_argument("--width", type=float, default=8.0, help="Figure canvas width in inches before tight export cropping. Example: --width 8, which sets an eight-inch canvas for consistent panel sizing.")
    parser.add_argument("--height", type=float, default=5.2, help="Figure canvas height in inches before tight export cropping. Example: --height 5.2, which sets a 5.2-inch canvas without changing data units.")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacement of generated files; otherwise existing outputs cause an error. Example: --overwrite, which regenerates figures and the manifest without modifying the workbook.")
    add_plot_style_argument(parser)
    add_color_style_argument(parser)
    parser.set_defaults(plot_style="publication-bold")


def run_main(command, args):
    """Validate the entire workbook, render its groups, and record source rows."""
    for name in ("width", "height"):
        value = _number(getattr(args, name), f"--{name}")
        if value <= 0:
            raise ValueError(f"--{name} must be positive.")
    figures = read_figures(args.input)
    color_style = resolve_color_style(getattr(args, "color_style", None))
    root = Path(args.output) if args.output is not None else Path(args.input).resolve().parent / "excel_figures"
    formats = list(dict.fromkeys(args.formats))
    outputs = [root / figure["sheet"] / f"figure_{figure['figure']:03d}.{extension}"
               for figure in figures for extension in formats]
    manifest = root / "excel_figures.json"
    if not args.overwrite:
        for path in [*outputs, manifest]:
            if io_path(path).exists():
                raise FileExistsError(f"Output exists: {path}. Use --overwrite to replace it.")
    for figure in figures:
        payload = figure_payload(figure, plot_style=args.plot_style, figsize=(args.width, args.height))
        payload["color_style"] = color_style
        for extension in formats:
            destination = root / figure["sheet"] / f"figure_{figure['figure']:03d}.{extension}"
            render_plot({**payload, "save": str(io_path(destination))})
    io_path(manifest).write_text(json.dumps({"input": str(Path(args.input).resolve()),
        "plot_style": args.plot_style, "color_style": color_style, "figsize_inches": [args.width, args.height],
        "outputs": [str(path) for path in outputs], "figures": figures}, indent=2, default=str), encoding="utf-8")
    print(f"Generated {len(figures)} figures ({len(outputs)} files) in {root}.")
    return 0
