from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from reaxkit.workflows.force_field_opt.plot_from_excel import (
    DATA_HEADERS, SETTINGS_HEADERS, SHEETS, figure_payload, read_figures, run_main,
)


def workbook_file(tmp_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEETS:
        sheet = workbook.create_sheet(name)
        sheet.append([*DATA_HEADERS, None, *SETTINGS_HEADERS])
        sheet.append([2, 10, "second", 20, -3, -4, None, 2, "Comparison", "Volume (A^3)", "Energy (kcal/mol)"])
        sheet.append([2, None, "first", 10, 0, 0])
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    return workbook, path


def test_grouping_sorting_and_original_values(tmp_path):
    workbook, path = workbook_file(tmp_path)
    sheet = workbook["bar_plots"]
    sheet.append([1, "report expression", "other", None, 7, 8, None, 1, "Other", "Structure", "Energy"])
    workbook.save(path)
    figures = read_figures(path)
    assert [(item["sheet"], item["figure"]) for item in figures] == [
        ("bar_plots", 1), ("bar_plots", 2), ("line_plots", 2), ("EOS_plots", 2)]
    assert [row["identifier"] for row in figures[1]["rows"]] == ["second", "first"]
    assert [row["x"] for row in figures[3]["rows"]] == [10, 20]
    assert [row["reaxff"] for row in figures[3]["rows"]] == [0, -3]


@pytest.mark.parametrize("cell,value,message", [
    ("A2", 1.5, "positive integer"), ("C2", None, "Identifier"),
    ("D2", None, "EOS requires"), ("E2", "bad", "finite number"),
    ("E2", "=1+1", "cached result"), ("F2", "#DIV/0!", "Excel error"),
    ("H3", 2, "duplicate settings"), ("J2", None, "axis labels"),
    ("H2", 1, "settings figure numbers"),
])
def test_invalid_workbook_rejected_before_rendering(tmp_path, cell, value, message):
    workbook, path = workbook_file(tmp_path)
    workbook["EOS_plots"][cell] = value
    workbook.save(path)
    with pytest.raises(ValueError, match=message):
        read_figures(path)


def test_categorical_lines_and_mixed_coordinates(tmp_path):
    workbook, path = workbook_file(tmp_path)
    sheet = workbook["line_plots"]
    sheet["D2"] = None
    workbook.save(path)
    with pytest.raises(ValueError, match="every point"):
        read_figures(path)
    sheet["D3"] = None
    workbook.save(path)
    figure = next(item for item in read_figures(path) if item["sheet"] == "line_plots")
    payload = figure_payload(figure, plot_style="publication-bold", figsize=(8, 5))
    assert payload["xticklabels"] == ["second", "first"]
    assert payload["series"][0]["x"] == [0, 1]


def test_render_and_overwrite_guard(tmp_path):
    _, path = workbook_file(tmp_path)
    args = SimpleNamespace(input=str(path), output=str(tmp_path / "figures"),
        formats=["png", "svg"], width=6, height=4, overwrite=False, plot_style="publication-bold")
    assert run_main("plot-from-excel", args) == 0
    outputs = list(Path(args.output).rglob("figure_*"))
    assert len(outputs) == 6 and all(path.stat().st_size > 1000 for path in outputs)
    assert "#0057b8" in (Path(args.output) / "bar_plots/figure_002.svg").read_text().lower()
    assert "#d62728" in (Path(args.output) / "bar_plots/figure_002.svg").read_text().lower()
    with pytest.raises(FileExistsError):
        run_main("plot-from-excel", args)


def test_cli_registration():
    from reaxkit.cli.main import build_parser
    args = build_parser("plot-from-excel").parse_args(["plot-from-excel", "--input", "input.xlsx"])
    assert args.plot_style == "publication-bold"
    assert args.formats == ["png", "svg"]


@pytest.mark.parametrize("output", [None, "custom_figures"])
def test_output_location_from_another_directory(tmp_path, monkeypatch, output):
    import json
    from reaxkit.cli.main import build_parser

    input_directory = tmp_path / "inputs"
    input_directory.mkdir()
    _, path = workbook_file(input_directory)
    run_directory = tmp_path / "run"
    run_directory.mkdir()
    monkeypatch.chdir(run_directory)
    command = ["plot-from-excel", "--input", "../inputs/input.xlsx", "--formats", "png"]
    if output is not None:
        command.extend(["--output", output])
    args = build_parser("plot-from-excel").parse_args(command)
    assert run_main("plot-from-excel", args) == 0
    destination = input_directory / "excel_figures" if output is None else run_directory / output
    manifest = json.loads((destination / "excel_figures.json").read_text())
    assert Path(manifest["input"]) == path
    assert len(list(destination.rglob("figure_*.png"))) == 3
    assert all(Path(filename).is_file() for filename in manifest["outputs"])
    assert not (run_directory / "excel_figures").exists()


@pytest.mark.parametrize("count", [2, 5])
def test_long_bar_labels_leave_more_plot_height(count, monkeypatch):
    import matplotlib.pyplot as plt
    from reaxkit.presentation.plot.renderers import grouped_bar

    monkeypatch.setattr(grouped_bar, "save_or_show", lambda figure, cfg: figure)
    figure = {"sheet": "bar_plots", "categorical": False, "title": "Comparison",
              "xlabel": "Structure", "ylabel": "Energy",
              "rows": [{"identifier": f"Long_structure_identifier_{index}", "x": None,
                        "reaxff": index + 1, "qm": index + 2} for index in range(count)]}
    payload = figure_payload(figure, plot_style="publication-bold", figsize=(8, 5.2))
    tilted = grouped_bar.GroupedBarPlotRenderer().render(payload)
    vertical = grouped_bar.GroupedBarPlotRenderer().render(
        {**payload, "label_rotation": 90, "label_horizontal_alignment": "center"})
    try:
        assert all(label.get_rotation() == 45 for label in tilted.axes[0].get_xticklabels())
        assert tilted.axes[0].get_position().height > vertical.axes[0].get_position().height
    finally:
        plt.close(tilted)
        plt.close(vertical)


def test_bare_command_shows_help(monkeypatch, capsys):
    import sys
    from reaxkit.cli.main import main

    monkeypatch.setattr(sys, "argv", ["reaxkit", "plot-from-excel"])
    with pytest.raises(SystemExit) as result:
        main(announce=False)
    assert result.value.code == 0
    output = capsys.readouterr().out
    assert "Examples:" in output
    assert "Scope:" in output
    assert "--input PtZn_training_figures.xlsx" in output


def test_help_follows_workflow_documentation_rules():
    import argparse
    from reaxkit.workflows.force_field_opt.plot_from_excel import build_parser

    parser = argparse.ArgumentParser()
    build_parser(parser)
    assert parser.formatter_class is argparse.RawTextHelpFormatter
    assert "Examples:" in parser.description and "Scope:" in parser.description
    for action in parser._actions:
        if action.dest == "help":
            continue
        assert action.help.count("Example:") == 1
        assert ", which " in action.help


def test_packaged_template_is_readable():
    from reaxkit.workflows.force_field_opt.get_ffield_opt_plots import _figure_generator_template_source
    figures = read_figures(_figure_generator_template_source())
    assert [figure["sheet"] for figure in figures] == list(SHEETS)


def test_missing_settings_empty_data_and_nonfinite_values(tmp_path):
    workbook, path = workbook_file(tmp_path)
    workbook["bar_plots"]["F2"] = "NaN"
    workbook.save(path)
    with pytest.raises(ValueError, match="finite number"):
        read_figures(path)
    for sheet in workbook:
        sheet.delete_rows(2, sheet.max_row)
    workbook.save(path)
    with pytest.raises(ValueError, match="no figure data"):
        read_figures(path)
